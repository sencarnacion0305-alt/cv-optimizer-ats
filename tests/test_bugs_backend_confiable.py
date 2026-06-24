"""
Bugs críticos del backend (base de análisis confiable).

Cubre los 9 casos requeridos:
  1-2. limpiar-vacante conserva requisitos/responsabilidades (EN y ES).
  3.   detectar "Backend Developer" como cargo objetivo.
  4.   titulo_cubierto NO es true cuando no hay cargo en la vacante.
  5.   parser detecta "New York, NY" como ubicación.
  6.   parser extrae titulo/empresa/periodo correctamente.
  7.   parser no usa una viñeta como empresa.
  8.   optimizar-cv devuelve score_antes válido (texto y DOCX) o explicación clara.
  9.   mejorar-bullets no inventa métricas reales (usa placeholders + advertencia).
"""
import io
import re

import pytest
from docx import Document

from app.models.schemas import AdaptarCVRequest
from app.services.adaptador import adaptar_cv, _detectar_titulo_vacante
from app.services.limpiador import limpiar_vacante
from app.services.mejorador_bullets import mejorar_bullets
from app.services.parser_ats import simular_parsing


VAC_EN = ("Backend Developer - We offer benefits, culture, snacks. "
          "Requirements: Python, FastAPI, PostgreSQL, Docker, AWS, CI/CD. "
          "Responsibilities: build REST APIs and write tests.")
VAC_ES = ("Desarrollador Backend. Beneficios: seguro medico y cultura flexible. "
          "Requisitos: Python, FastAPI, PostgreSQL, Docker, AWS, CI/CD. "
          "Responsabilidades: construir APIs REST, escribir pruebas y optimizar consultas SQL.")

CV_DEMO = """Alex Rivera
Email: alex.rivera@example.com | Phone: +1 555 0100 | LinkedIn: linkedin.com/in/alexrivera | New York, NY

SUMMARY
Backend Developer with 4 years of experience building REST APIs with Python, FastAPI, PostgreSQL, Docker and CI/CD.

EXPERIENCE
Backend Developer - Nova Systems | January 2021 - May 2025
- Developed REST APIs with Python and FastAPI for internal platforms.
- Optimized PostgreSQL queries and reduced response time by 35%.
- Implemented Docker containers and CI/CD workflows.

EDUCATION
B.S. Information Systems - Demo University, 2018

SKILLS
Python, FastAPI, PostgreSQL, Docker, Git, REST APIs, Linux, CI/CD, testing."""


# ── Bug 1: limpiar-vacante EN nunca vacía y conserva requisitos ──
def test_limpiar_vacante_en_conserva_requisitos():
    out = limpiar_vacante(VAC_EN)
    low = out.lower()
    assert out.strip()                       # nunca vacío
    assert "requirements" in low
    assert "responsibilities" in low
    assert "python" in low and "docker" in low
    # ruido eliminado
    assert "we offer" not in low
    assert "benefit" not in low
    assert "snacks" not in low


# ── Bug 1: limpiar-vacante ES nunca vacía y conserva requisitos ──
def test_limpiar_vacante_es_conserva_requisitos():
    out = limpiar_vacante(VAC_ES)
    low = out.lower()
    assert out.strip()
    assert "requisitos" in low
    assert "responsabilidades" in low
    assert "python" in low and "docker" in low
    assert "beneficios" not in low
    assert "seguro medico" not in low


# ── Bug 2: detectar "Backend Developer" como cargo (EN y ES) ──
def test_detecta_backend_developer():
    assert _detectar_titulo_vacante(
        "Backend Developer Python FastAPI PostgreSQL Docker AWS CI/CD") == "Backend Developer"


@pytest.mark.parametrize("vac,esperado", [
    ("Software Engineer Python", "Software Engineer"),
    ("Ingeniero de Software con experiencia", "Ingeniero de Software"),
    ("Data Analyst SQL Power BI", "Data Analyst"),
    ("Gerente de Proyecto", "Gerente de Proyecto"),
    ("QA Engineer Selenium", "QA Engineer"),
    ("Ingeniero DevOps Kubernetes", "Ingeniero DevOps"),
])
def test_detecta_cargos_en_es(vac, esperado):
    assert _detectar_titulo_vacante(vac) == esperado


# ── Bug 2: titulo_cubierto NO es true si no hay cargo en la vacante ──
def test_titulo_no_cubierto_si_vacante_sin_cargo():
    req = AdaptarCVRequest(cv_texto="Alex Rivera Python Docker",
                           vacante_texto="Buscamos a alguien con muchas ganas de aprender.")
    resp = adaptar_cv(req)
    assert resp.titulo_vacante is None
    assert resp.titulo_cubierto is False


def test_titulo_cubierto_true_cuando_aparece_en_cv():
    req = AdaptarCVRequest(
        cv_texto="Alex Rivera Backend Developer Python FastAPI Docker REST APIs",
        vacante_texto="Backend Developer Python FastAPI PostgreSQL Docker")
    resp = adaptar_cv(req)
    assert resp.titulo_vacante == "Backend Developer"
    assert resp.titulo_cubierto is True


# ── Bug 3: parser ubicación / titulo / empresa / periodo ──
def test_parser_detecta_ubicacion_ciudad_estado():
    r = simular_parsing(CV_DEMO)
    assert r["campos"]["ubicacion"] == "New York, NY"


def test_parser_extrae_titulo_empresa_periodo():
    r = simular_parsing(CV_DEMO)
    p = r["puestos"][0]
    assert p["titulo"] == "Backend Developer"
    assert p["empresa"] == "Nova Systems"
    assert p["periodo"] == "January 2021 - May 2025"


def test_parser_no_usa_bullet_como_empresa():
    r = simular_parsing(CV_DEMO)
    for p in r["puestos"]:
        assert not p["empresa"].lstrip().startswith("-")
        assert "Developed REST APIs" not in p["empresa"]


def test_parser_detecta_contacto_educacion_skills():
    r = simular_parsing(CV_DEMO)
    c = r["campos"]
    assert c["nombre"] == "Alex Rivera"
    assert c["email"] == "alex.rivera@example.com"
    assert c["telefono"] == "+1 555 0100"
    assert c["linkedin"] == "linkedin.com/in/alexrivera"
    assert r["educacion"]                       # educación detectada
    assert "python" in [s.lower() for s in r["skills"]]


@pytest.mark.parametrize("linea,periodo_ok", [
    ("Backend Developer - Nova | January 2021 - May 2025", "January 2021 - May 2025"),
    ("Backend Developer - Nova | Jan 2021 - May 2025", "Jan 2021 - May 2025"),
    ("Desarrollador Backend - Nova | Enero 2021 - Mayo 2025", "Enero 2021 - Mayo 2025"),
    ("Backend Developer - Nova | 2021 - 2025", "2021 - 2025"),
    ("Backend Developer - Nova | 2021 - Present", "2021 - Present"),
])
def test_parser_formatos_de_fecha(linea, periodo_ok):
    from app.services.parser_ats import _detectar_bloques_fechados
    b = _detectar_bloques_fechados([linea])
    assert b and b[0]["periodo"] == periodo_ok
    assert b[0]["titulo"] in ("Backend Developer", "Desarrollador Backend")
    assert b[0]["empresa"] == "Nova"


# ── Bug 5: mejorar-bullets no inventa métricas reales ──
def test_bullets_metricas_son_placeholders_con_advertencia():
    cv = ("Experience\nBackend Developer\n"
          "- Implemented Docker containers and CI/CD workflows\n"
          "- Worked on REST APIs with Python")
    r = mejorar_bullets(cv)
    assert r["advertencia"]                     # advertencia presente
    assert "reales" in r["advertencia"].lower()
    for m in r["mejoras"]:
        if m["metrica_agregada"]:
            # marcador editable presente
            assert ("[number]" in m["mejorado"]
                    or "[% estimado]" in m["mejorado"]
                    or "[$ amount]" in m["mejorado"])
            # NINGÚN dígito crudo fuera de los corchetes (nada inventado como real)
            sin_placeholder = re.sub(r"\[[^\]]*\]", "", m["mejorado"])
            assert not re.search(r"\d", sin_placeholder)


# ── Bug 4: optimizar-cv score_antes válido (texto y DOCX) o explicación clara ──
def _client():
    from fastapi.testclient import TestClient
    from app.main import app
    return TestClient(app)


def test_optimizar_cv_texto_score_antes_valido():
    r = _client().post("/api/v1/optimizar-cv",
                       data={"cv_texto": CV_DEMO, "vacante_texto": "Backend Developer Python"})
    assert r.status_code == 200
    j = r.json()
    assert isinstance(j["score_antes"], int) and 0 <= j["score_antes"] <= 100
    assert j["archivo_base64"]                  # CV descargable


def test_optimizar_cv_docx_score_antes_valido():
    doc = Document()
    for line in CV_DEMO.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    r = _client().post(
        "/api/v1/optimizar-cv",
        files={"archivo": ("cv.docx", buf.read(),
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"vacante_texto": "Backend Developer Python"})
    assert r.status_code == 200
    j = r.json()
    # El bug era score_antes=null para entrada por archivo (doble lectura del upload).
    assert isinstance(j["score_antes"], int) and 0 <= j["score_antes"] <= 100
    assert j["archivo_base64"]


def test_optimizar_cv_sin_cv_explica_claramente():
    r = _client().post("/api/v1/optimizar-cv", data={"vacante_texto": "x"})
    assert r.status_code == 400
    assert "CV" in r.json()["detail"]           # no falla en silencio
